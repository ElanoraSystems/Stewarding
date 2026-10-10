import csv
import io

import openpyxl
from sqlalchemy import select

from app.models import AuditLog, ChecklistInstance

from .helpers import MGR, STEWARD, STEWARD_RB, SUP, SUP_RB, balance, item_id, loc_id


def generate(client, login, db, code="MK"):
    client.post("/api/checklists/generate", json={"location_id": loc_id(db, code)}, headers=login(MGR))


def rep(client, h, period="daily", **params):
    query = "&".join(f"{k}={v}" for k, v in params.items())
    r = client.get(f"/api/reports/{period}?{query}", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def sec(report, sid):
    return next(s for s in report["sections"] if s["id"] == sid)


def totals(section):
    return {t["label"]: t["value"] for t in section["totals"]}


def complete_tasks(client, h, n, skip=()):
    done = 0
    for t in client.get("/api/checklists/instances?size=200", headers=h).json()["items"]:
        if t["frequency"] != "weekly" and not any(s in t["description"] for s in skip) and done < n:
            client.post(f"/api/checklists/instances/{t['id']}/complete", json={}, headers=h)
            done += 1


def test_checklist_report_shows_numerator_denominator_and_reconciles_with_records(client, login, db, clock):
    generate(client, login, db)
    complete_tasks(client, login(STEWARD), 3, skip=("filters", "next shift"))
    clock.set(2026, 3, 2, 10, 0)  # 14:00 Dubai, after the 13:30 deadline
    report = rep(client, login(MGR), anchor="2026-03-02", location_id=loc_id(db, "MK"))
    s = sec(report, "checklists")
    t = totals(s)
    assert t["Compliance"] == "3/14 (21.4%)" and t["Completed"] == 3 and t["Overdue"] == 11
    assert t["Not yet due (excluded)"] == 10  # weekly tasks are due at 17:00
    all_rows = len(db.scalars(select(ChecklistInstance).where(ChecklistInstance.location_id == loc_id(db, "MK"))).all())
    assert len(s["rows"]) == all_rows == 24
    assert sum(t[k] for k in ("Completed", "Verified", "Pending verification", "Waived (excluded)", "Overdue", "Failed", "Not yet due (excluded)")) == 24


def test_pending_verification_counts_as_done_and_waived_is_excluded(client, login, db, clock):
    generate(client, login, db)
    h, sup = login(STEWARD), login(SUP)
    items = client.get("/api/checklists/instances?size=200", headers=h).json()["items"]
    filters = next(i for i in items if "dishwasher filters" in i["description"])
    client.post(f"/api/checklists/instances/{filters['id']}/complete", json={}, headers=h)
    waived = next(i for i in items if "Clean walls" in i["description"])
    client.post(f"/api/checklists/instances/{waived['id']}/waive", json={"reason": "Wall under repair"}, headers=sup)
    clock.set(2026, 3, 2, 10, 0)
    t = totals(sec(rep(client, login(MGR), anchor="2026-03-02", location_id=loc_id(db, "MK")), "checklists"))
    assert t["Pending verification"] == 1 and t["Waived (excluded)"] == 1
    assert t["Compliance"] == "1/13 (7.7%)"


def test_attendance_report_scheduled_present_absent_and_unscheduled(client, login, db, clock):
    mk = loc_id(db, "MK")
    s = sec(rep(client, login(MGR), anchor="2026-03-02", location_id=mk), "attendance")
    assert totals(s)["Scheduled"] == 1 and totals(s)["Absent"] == 0 and s["rows"][0]["status"] == "not_yet_due"
    clock.set(2026, 3, 2, 2, 20)  # 06:20, past the grace period
    s = sec(rep(client, login(MGR), anchor="2026-03-02", location_id=mk), "attendance")
    assert totals(s)["Absent"] == 1 and s["rows"][0]["status"] == "absent"
    client.post("/api/attendance/check-in", json={"location_id": mk}, headers=login(STEWARD))
    client.post("/api/attendance/check-in", json={"location_id": mk}, headers=login(MGR))
    s = sec(rep(client, login(MGR), anchor="2026-03-02", location_id=mk), "attendance")
    t = totals(s)
    assert t["Present"].startswith("1/1") and t["Late arrivals"] == 1 and t["Absent"] == 0
    assert {r["status"] for r in s["rows"]} == {"on_duty", "unscheduled"}


def test_weekly_trend_never_counts_future_days(client, login, db, clock):
    clock.set(2026, 3, 4, 8, 0)  # Wednesday afternoon
    trend = sec(rep(client, login(MGR), "weekly", anchor="2026-03-04", location_id=loc_id(db, "MK")), "attendance_trend")
    assert [r["date"] for r in trend["rows"]] == ["2026-03-02", "2026-03-03", "2026-03-04"]
    assert [r["absent"] for r in trend["rows"]] == [1, 1, 1]  # nobody checked in on any past or current day


def test_temperature_compliance_counts_only_checks_due_so_far(client, login, db, clock):
    clock.set(2026, 3, 2, 2, 45)  # 06:45: wash and first final-rinse checks are due, 30 minute grace has passed
    h = login(STEWARD)
    eq = next(e for e in client.get("/api/temperature/equipment", headers=h).json() if e["code"] == "DW-MK-1")
    wash = next(s for s in eq["schedules"] if s["stage"] == "wash")
    client.post("/api/temperature/readings", headers=h, json={"equipment_id": eq["id"], "schedule_id": wash["id"], "temp_c": 60})
    s = sec(rep(client, login(MGR), anchor="2026-03-02", location_id=loc_id(db, "MK")), "temperature")
    assert totals(s)["Readings"] == 1
    weekly = sec(rep(client, login(MGR), "weekly", anchor="2026-03-02", location_id=loc_id(db, "MK")), "temperature_compliance")
    dw = next(r for r in weekly["rows"] if r["equipment"] == "Main Kitchen Dishwasher 1")
    assert (dw["done"], dw["due"], dw["missed"], dw["percent"]) == (1, 2, 1, 50.0)
    gw = next(r for r in weekly["rows"] if "Glasswasher" in r["equipment"])
    assert (gw["done"], gw["due"], gw["missed"]) == (0, 2, 2)
    assert totals(weekly)["Overall"] == "1/4 (25.0%)"


def test_failed_readings_appear_in_monthly_corrective_action_history(client, login, db, clock):
    h = login(STEWARD)
    eq = next(e for e in client.get("/api/temperature/equipment", headers=h).json() if e["code"] == "DW-MK-1")
    wash = next(s for s in eq["schedules"] if s["stage"] == "wash")
    client.post("/api/temperature/readings", headers=h, json={
        "equipment_id": eq["id"], "schedule_id": wash["id"], "temp_c": 40, "corrective_action": "Descaled"})
    s = sec(rep(client, login(MGR), "monthly", anchor="2026-03-15"), "corrective_actions")
    assert totals(s) == {"Failed readings": 1, "Resolved": 0}
    assert s["rows"][0]["status"] == "retest_due" and s["rows"][0]["action"] == "Descaled"
    fails = sec(rep(client, login(MGR), "weekly", anchor="2026-03-02"), "equipment_failures")
    assert fails["rows"][0]["failed"] == 1 and fails["rows"][0]["location"] == "Main Kitchen"


def test_breakage_totals_match_incidents(client, login, db, clock):
    h = login(STEWARD)
    for qty, code in ((2, "PL-001"), (3, "BW-001")):
        client.post("/api/breakage", headers=h, json={"location_id": loc_id(db, "MK"), "item_id": item_id(db, code),
                                                       "qty": qty, "incident_type": "breakage"})
    rejected = client.post("/api/breakage", headers=h, json={"location_id": loc_id(db, "MK"), "item_id": item_id(db, "PL-001"),
                                                              "qty": 9, "incident_type": "damage"}).json()
    client.post(f"/api/breakage/{rejected['id']}/reject", json={"reason": "Duplicate report"}, headers=login(SUP))
    daily = sec(rep(client, login(MGR), anchor="2026-03-02", location_id=loc_id(db, "MK")), "breakage")
    assert len(daily["rows"]) == 3
    assert totals(daily) == {"Incidents (excluding rejected)": 2, "Units": 5, "Estimated cost": 11.3}
    summary = sec(rep(client, login(MGR), "weekly", anchor="2026-03-02"), "breakage_summary")
    assert {r["item"]: r["qty"] for r in summary["rows"]} == {"Dinner Plate 27cm": 2, "Soup Bowl": 3}


def test_monthly_reconciliation_closing_matches_live_balance(client, login, db, clock):
    base = {"item_id": item_id(db, "PL-001"), "location_id": loc_id(db, "MK")}
    client.post("/api/inventory/transactions", json={**base, "txn_type": "receipt", "qty": 30}, headers=login(SUP))
    s = sec(rep(client, login(MGR), "monthly", anchor="2026-03-10", location_id=loc_id(db, "MK")), "reconciliation")
    row_ = next(r for r in s["rows"] if r["item"].startswith("PL-001"))
    assert row_["closing"] == balance(db, "PL-001", "MK") == 230 and row_["in"] == 230 and row_["out"] == 0
    consumption = sec(rep(client, login(MGR), "monthly", anchor="2026-03-10"), "stock_consumption")
    assert any(r["item"] == "Dinner Plate 27cm" and r["received"] == 30 for r in consumption["rows"])
    replen = sec(rep(client, login(MGR), "monthly", anchor="2026-03-10"), "replenishment")
    assert any(r["location"] == "Rooftop Bar" and r["item"] == "Dinner Plate 27cm" and r["order"] == 260 for r in replen["rows"])


def test_scorecard_and_organization_versus_location_reports(client, login, db, clock):
    org = sec(rep(client, login(MGR), "monthly", anchor="2026-03-10"), "scorecard")
    assert {r["location"] for r in org["rows"]} == {"Main Kitchen", "Banquet Hall", "Rooftop Bar"}
    one = rep(client, login(MGR), "monthly", anchor="2026-03-10", location_id=loc_id(db, "RB"))
    assert one["locations"] == ["Rooftop Bar"] and [r["location"] for r in sec(one, "scorecard")["rows"]] == ["Rooftop Bar"]
    sup = rep(client, login(SUP), "monthly", anchor="2026-03-10")
    assert set(sup["locations"]) == {"Main Kitchen", "Banquet Hall"}
    assert client.get(f"/api/reports/monthly?location_id={loc_id(db, 'RB')}", headers=login(SUP)).status_code == 403


def test_period_labels_timestamps_and_validation(client, login, db, clock):
    m = login(MGR)
    week = rep(client, m, "weekly", anchor="2026-03-04")
    assert week["period"] == {"type": "weekly", "label": "Week 2026-03-02 to 2026-03-08", "from": "2026-03-02", "to": "2026-03-08"}
    month = rep(client, m, "monthly", anchor="2026-03-04")
    assert month["period"]["label"] == "March 2026" and month["period"]["to"] == "2026-03-31"
    assert month["generated_at"] == "2026-03-02T02:05:00Z" and month["timezones"] == ["Asia/Dubai", "Europe/London"]
    assert client.get("/api/reports/yearly", headers=m).status_code == 404
    assert client.get("/api/reports", headers=m).json()["daily"][0] == "attendance"
    assert client.get("/api/reports/daily", headers=login(STEWARD)).status_code == 403


def test_report_dates_follow_the_location_timezone(client, login, db, clock):
    clock.set(2026, 3, 1, 21, 30)  # still Sunday in UTC, already Monday 01:30 in Dubai
    client.post("/api/attendance/check-in", json={"location_id": loc_id(db, "MK")}, headers=login(STEWARD))
    monday = sec(rep(client, login(MGR), anchor="2026-03-02", location_id=loc_id(db, "MK")), "attendance")
    assert [r["employee"] for r in monday["rows"] if r["status"] != "not_yet_due"] == ["Sid Steward"]
    sunday = sec(rep(client, login(MGR), anchor="2026-03-01", location_id=loc_id(db, "MK")), "attendance")
    assert not [r for r in sunday["rows"] if r["employee"] == "Sid Steward" and r["check_in"]]


def test_employee_filter(client, login, db, clock):
    clock.set(2026, 3, 2, 10, 0)
    emps = {e["employee_code"]: e["id"] for e in client.get("/api/employees?size=50", headers=login(MGR)).json()["items"]}
    s = sec(rep(client, login(MGR), "monthly", anchor="2026-03-02", employee_id=emps["E004"]), "attendance_hours")
    assert [r["employee"] for r in s["rows"]] == ["Sid Steward"]


def export(client, h, fmt, period="daily", **params):
    query = "&".join(f"{k}={v}" for k, v in {"fmt": fmt, "anchor": "2026-03-02", **params}.items())
    return client.get(f"/api/reports/{period}/export?{query}", headers=h)


def test_export_is_manager_only_and_audited(client, login, db, clock):
    for who in (SUP, STEWARD):
        assert export(client, login(who), "csv").status_code == 403
    r = export(client, login(MGR), "csv")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert db.scalar(select(AuditLog.id).where(AuditLog.action == "report.export"))
    assert export(client, login(MGR), "docx").status_code == 422
    assert client.get("/api/reports/yearly/export", headers=login(MGR)).status_code == 404


def test_csv_export_has_period_timestamp_and_totals(client, login, db, clock):
    generate(client, login, db)
    text = export(client, login(MGR), "csv", location_id=loc_id(db, "MK")).content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    flat = [c for r in rows for c in r]
    assert ["Period", "2026-03-02 (2026-03-02 to 2026-03-02)"] in rows
    assert ["Generated at (UTC)", "2026-03-02T02:05:00Z"] in rows
    assert "Compliance" in flat and "Employee attendance and shift coverage" in flat


def test_xlsx_and_pdf_exports_are_valid_files(client, login, db, clock):
    generate(client, login, db)
    xlsx = export(client, login(MGR), "xlsx", "weekly")
    assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats")
    book = openpyxl.load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames[0] == "Report" and len(book.sheetnames) == 8
    assert book["Report"]["A5"].value == "Generated at (UTC)"
    pdf = export(client, login(MGR), "pdf", "monthly")
    assert pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")


def test_exports_neutralise_spreadsheet_formulas(client, login, db, clock):
    emps = {e["employee_code"]: e["id"] for e in client.get("/api/employees?size=50", headers=login(MGR)).json()["items"]}
    client.put(f"/api/employees/{emps['E004']}", json={"full_name": "=HYPERLINK(\"http://x\",\"a\")"}, headers=login(MGR))
    text = export(client, login(MGR), "csv", location_id=loc_id(db, "MK")).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text and ",=HYPERLINK" not in text
    book = openpyxl.load_workbook(io.BytesIO(export(client, login(MGR), "xlsx", location_id=loc_id(db, "MK")).content))
    values = [c.value for sheet in book for row in sheet.iter_rows() for c in row if isinstance(c.value, str)]
    assert any(v.startswith("'=HYPERLINK") for v in values) and not any(v.startswith("=") for v in values)


def test_other_location_staff_cannot_run_reports_for_mk(client, login, db, clock):
    assert client.get(f"/api/reports/daily?location_id={loc_id(db, 'MK')}", headers=login(SUP_RB)).status_code == 403
    assert client.get("/api/reports/daily", headers=login(STEWARD_RB)).status_code == 403
