import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.utils import ImageReader
from reportlab.platypus import Flowable, Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .organization import stamp

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe(value):
    """Stops spreadsheet apps from running text that starts like a formula."""
    if isinstance(value, str) and value[:1] in FORMULA_PREFIXES:
        return "'" + value
    return value


def _org_lines(org: dict | None) -> list[str]:
    """The organization's name and details, one line each, for the head of a report."""
    if not org or not org.get("name"):
        return []
    contact = " | ".join(part for part in (org.get("phone"), org.get("email"), org.get("website")) if part)
    legal = org.get("legal_name") or ""
    registration = f"Reg. {org['registration_number']}" if org.get("registration_number") else ""
    return [line for line in (org["name"], legal if legal != org["name"] else "", org.get("address") or "", contact, registration) if line]


def _pdf_masthead(org: dict | None, styles, width: float) -> list:
    """Logo on the left, organization name and details on the right; nothing when no organization is set."""
    lines = _org_lines(org)
    if not lines:
        return []
    text = [Paragraph(f"<b>{_esc(lines[0])}</b>", styles["Heading2"])] + [Paragraph(_esc(line), styles["BodyText"].clone("org", fontSize=8, leading=10)) for line in lines[1:]]
    logo = None
    if org.get("logo_path"):
        w, h = ImageReader(org["logo_path"]).getSize()
        height = 42
        logo = Image(org["logo_path"], width=height * w / h, height=height)
    table = Table([[logo or "", text]], colWidths=[(logo.drawWidth + 12) if logo else 0.1, width - ((logo.drawWidth + 12) if logo else 0.1)])
    table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#9ca3af")),
                               ("BOTTOMPADDING", (0, 0), (-1, 0), 6), ("LEFTPADDING", (0, 0), (0, 0), 0)]))
    return [table, Spacer(1, 6)]


def _generated(report: dict) -> str:
    tz = (report.get("organization") or {}).get("timezone") or "Asia/Kuwait"
    return stamp(datetime.fromisoformat(report["generated_at"].rstrip("Z")), tz)


def _header(report: dict) -> list[tuple[str, str]]:
    period = report["period"]
    return [("Report", f"{period['type'].title()} report"), ("Period", f"{period['label']} ({period['from']} to {period['to']})"),
            ("Locations", ", ".join(report["locations"])), ("Location time zones", ", ".join(report["timezones"])),
            ("Generated", _generated(report)), ("Prepared by", report.get("prepared_by") or "-"),
            ("Position", report.get("prepared_position") or "-")]


def to_csv(report: dict) -> bytes:
    buf = io.StringIO()
    out = csv.writer(buf)
    for line in _org_lines(report.get("organization")):
        out.writerow([safe(line)])
    for label, value in _header(report):
        out.writerow([label, safe(value)])
    for sec in report["sections"]:
        out.writerow([])
        out.writerow([safe(sec["title"])])
        out.writerow([safe(c["label"]) for c in sec["columns"]])
        for row in sec["rows"]:
            out.writerow([safe(row.get(c["key"])) for c in sec["columns"]])
        for total in sec["totals"]:
            out.writerow([safe(total["label"]), safe(total["value"])])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(report: dict) -> bytes:
    book = Workbook()
    info = book.active
    info.title = "Report"
    org_lines = _org_lines(report.get("organization"))
    for line in org_lines:
        info.append([safe(line)])
    if org_lines:
        info["A1"].font = Font(bold=True, size=13)
    for label, value in _header(report):
        info.append([label, safe(value)])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 70
    used = {"Report"}
    for index, sec in enumerate(report["sections"], 1):
        name = "".join(ch for ch in sec["title"] if ch not in "[]:*?/\\")[:28] or f"Section {index}"
        while name in used:
            name = f"{name[:25]}_{index}"
        used.add(name)
        sheet = book.create_sheet(name)
        sheet.append([safe(c["label"]) for c in sec["columns"]])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in sec["rows"]:
            sheet.append([safe(row.get(c["key"])) for c in sec["columns"]])
        for total in sec["totals"]:
            sheet.append([])
            sheet.append([safe(total["label"]), safe(total["value"])])
        for i in range(len(sec["columns"])):
            sheet.column_dimensions[chr(65 + i % 26)].width = 20
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def to_pdf(report: dict) -> bytes:
    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=7, leading=8)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24,
                            title=f"{report['period']['type'].title()} report")
    story = _pdf_masthead(report.get("organization"), styles, landscape(A4)[0] - 48)
    story.append(Paragraph(f"{report['period']['type'].title()} operations report", styles["Title"]))
    story += [Paragraph(f"<b>{label}:</b> {_esc(value)}", styles["Normal"]) for label, value in _header(report)[1:]]
    for sec in report["sections"]:
        story += [Spacer(1, 12), Paragraph(_esc(sec["title"]), styles["Heading3"])]
        if not sec["rows"]:
            story.append(Paragraph("No records for this period.", styles["Normal"]))
        else:
            data = [[Paragraph(f"<b>{_esc(c['label'])}</b>", cell) for c in sec["columns"]]]
            data += [[Paragraph(_esc(row.get(c["key"])), cell) for c in sec["columns"]] for row in sec["rows"]]
            table = Table(data, repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef0")),
                                       ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(table)
        if sec["totals"]:
            story.append(Paragraph(" | ".join(f"<b>{_esc(t['label'])}:</b> {_esc(t['value'])}" for t in sec["totals"]), styles["Normal"]))
    doc.build(story, onFirstPage=_footer_for(report), onLaterPages=_footer_for(report))
    return buf.getvalue()


def _footer_for(report: dict, title: str | None = None):
    org = report.get("organization") or {}
    left = " - ".join(part for part in (org.get("name"), title or f"{report['period']['type'].title()} report", org.get("report_footer")) if part)

    def draw(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(24, 14, left[:170])
        canvas.drawRightString(landscape(A4)[0] - 24, 14, f"Page {document.page}")
        canvas.restoreState()

    return draw


def _esc(value) -> str:
    text = "" if value is None else str(value)
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


FILLS = {"ok": "DCFCE7", "pending": "FEF3C7", "fail": "FEE2E2"}
WIDE = {"corrective_action": 3, "task": 3, "comment": 2, "approved_by": 2, "check": 1.8, "verified_by": 1.5, "limits": 1.4,
        "measure": 1.5, "date": 1.3, "time": 0.8, "result": 0.9, "value": 1.1}


SYMBOLS = {"ok": "\u2713", "pending": "\u2713?", "fail": "\u2717", "waived": "W", "na": "-"}
COLORS = {"ok": "#15803d", "pending": "#b45309", "fail": "#b91c1c"}


class Mark(Flowable):
    """A tick or a cross drawn as lines, so it prints the same in every viewer."""

    def __init__(self, state: str):
        super().__init__()
        self.state, self.width, self.height = state, 14, 9

    def wrap(self, *_):
        return self.width, self.height

    def draw(self):
        canvas = self.canv
        canvas.setStrokeColor(colors.HexColor(COLORS[self.state]))
        canvas.setLineWidth(1.4)
        if self.state == "fail":
            canvas.line(4, 1, 10, 8)
            canvas.line(4, 8, 10, 1)
        else:
            path = canvas.beginPath()
            path.moveTo(3.5, 4.5)
            path.lineTo(6, 1.5)
            path.lineTo(11, 8)
            canvas.drawPath(path, stroke=1, fill=0)
            if self.state == "pending":
                canvas.setDash(1, 1)
                canvas.line(2, 0, 12, 0)


def _flat(value):
    """A chart cell as text for spreadsheets: a symbol plus the initials of who did it and who approved it."""
    if isinstance(value, dict):
        who = "/".join(part for part in (value.get("by"), value.get("ver")) if part)
        return f"{SYMBOLS[value['s']]} {who}".strip()
    return value


def _signature(width: float) -> Table:
    sign = Table([["Prepared by", "", "Reviewed and approved by", "", "Date", ""]], colWidths=[width * f for f in (.12, .22, .2, .22, .06, .18)],
                 rowHeights=28)
    sign.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8),
                              ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#e8eef0")),
                              ("BACKGROUND", (2, 0), (2, 0), colors.HexColor("#e8eef0")), ("BACKGROUND", (4, 0), (4, 0), colors.HexColor("#e8eef0"))]))
    return sign


def _mark(key: str, value, style) -> object:
    if isinstance(value, dict):
        state = value["s"]
        center = style.clone("center", alignment=1)
        if state in COLORS:
            who = "/".join(part for part in (value.get("by"), value.get("ver")) if part)
            return [Mark(state), Paragraph(_esc(who), center.clone("tiny", fontSize=4.5, leading=5))]
        return Paragraph(f'<font color="#6b7280">{SYMBOLS[state]}</font>', center)
    text = _esc(value)
    return Paragraph(f'<font color="#b91c1c"><b>{text}</b></font>' if key == "result" and value == "FAIL" else text, style)


def _record_meta(doc: dict) -> list[tuple[str, str]]:
    return [("Period", doc["period"]["label"]), ("Locations", ", ".join(doc["locations"]) or "-"),
            ("Prepared by", doc["prepared_by"]), ("Position", doc.get("prepared_position") or "-"), ("Generated", doc["generated_at"])]


def record_to_csv(doc: dict) -> bytes:
    buf = io.StringIO()
    out = csv.writer(buf)
    for line in _org_lines(doc.get("organization")):
        out.writerow([safe(line)])
    out.writerow([safe(doc["title"])])
    for label, value in _record_meta(doc):
        out.writerow([label, safe(value)])
    for group in doc["groups"]:
        out.writerow([])
        out.writerow([safe(group["title"])])
        out.writerow([safe(label) for _, label in group["columns"]])
        for row in group["rows"]:
            out.writerow([safe(_flat(row.get(key))) for key, _ in group["columns"]])
        for label, value in group["totals"]:
            out.writerow([label, value])
    return buf.getvalue().encode("utf-8-sig")


def record_to_xlsx(doc: dict) -> bytes:
    book = Workbook()
    book.remove(book.active)
    used: set[str] = set()
    for index, group in enumerate(doc["groups"] or [{"title": "No records", "meta": [], "columns": [], "rows": [], "totals": []}], 1):
        name = "".join(ch for ch in group["title"] if ch not in "[]:*?/\\")[:28] or f"Sheet {index}"
        while name in used:
            name = f"{name[:25]}_{index}"
        used.add(name)
        sheet = book.create_sheet(name)
        org_lines = _org_lines(doc.get("organization"))
        for line in org_lines:
            sheet.append([safe(line)])
        if org_lines:
            sheet["A1"].font = Font(bold=True, size=14)
        sheet.append([safe(doc["title"])])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True, size=13)
        sheet.append([safe(group["title"])])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
        for label, value in [*_record_meta(doc), *group["meta"]]:
            sheet.append([label, safe(value)])
        sheet.append([])
        header_row = sheet.max_row + 1
        sheet.append([safe(label) for _, label in group["columns"]])
        for cell in sheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EEF0")
        for row in group["rows"]:
            sheet.append([safe(_flat(row.get(key))) for key, _ in group["columns"]])
            if group.get("layout") == "chart":
                for cell in sheet[sheet.max_row]:
                    state = next((s for s, sym in SYMBOLS.items() if isinstance(cell.value, str) and cell.value.startswith(sym) and s != "na"), None)
                    if state in FILLS:
                        cell.fill = PatternFill("solid", fgColor=FILLS[state])
        sheet.append([])
        for label, value in group["totals"]:
            sheet.append([label, value])
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
        for i, (key, _) in enumerate(group["columns"]):
            letter = sheet.cell(row=1, column=i + 1).column_letter
            sheet.column_dimensions[letter].width = 9 if group.get("layout") == "chart" and key != "task" else 16 * WIDE.get(key, 1)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def record_to_pdf(doc: dict) -> bytes:
    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)
    head = cell.clone("head", fontName="Helvetica-Bold")
    title_style = styles["Title"].clone("record_title", fontSize=20, leading=24)
    buf = io.BytesIO()
    page = landscape(A4)
    width = page[0] - 48

    org = doc.get("organization") or {}
    note = " - ".join(part for part in (org.get("name"), f"{doc['title']} - {doc['period']['label']}", org.get("report_footer")) if part)

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(24, 14, note[:170])
        canvas.drawRightString(page[0] - 24, 14, f"Page {document.page}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(buf, pagesize=page, leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=30, title=doc["title"])
    story = _pdf_masthead(org, styles, width) + [Paragraph(_esc(doc["title"]), title_style)]
    info = Table([[Paragraph(f"<b>{_esc(label)}</b>", cell), Paragraph(_esc(value), cell)] for label, value in _record_meta(doc)],
                 colWidths=[90, width - 90])
    info.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#e8eef0"))]))
    story.append(info)
    if not doc["groups"]:
        story += [Spacer(1, 14), Paragraph("No records for this period.", styles["Normal"])]
    for number, group in enumerate(doc["groups"]):
        if group.get("page_break") and number:
            story.append(PageBreak())
        chart = group.get("layout") == "chart"
        weights = [(8 if key == "task" else 1) if chart else WIDE.get(key, 1) for key, _ in group["columns"]]
        widths = [width * w / sum(weights) for w in weights]
        story += [Spacer(1, 14), Paragraph(_esc(group["title"]), styles["Heading3"]),
                  Paragraph(" | ".join(f"<b>{_esc(label)}:</b> {_esc(value)}" for label, value in group["meta"]), styles["Normal"]),
                  Spacer(1, 4)]
        if group.get("summary_first"):
            box = Table([[Paragraph(_esc(label), cell) for label, _ in group["totals"]],
                         [Paragraph(f"<b>{_esc(value)}</b>", cell.clone("big", fontSize=13, leading=15)) for _, value in group["totals"]]],
                        colWidths=[width / len(group["totals"])] * len(group["totals"]))
            box.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef0")),
                                     ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
            story += [box, Spacer(1, 8)]
        if group.get("legend"):
            story.append(Paragraph(_esc(group["legend"]), cell))
        data = [[Paragraph(_esc(label).replace("\n", "<br/>"), head if not chart else head.clone("chart_head", alignment=1 if label != "Task" else 0)) for _, label in group["columns"]]]
        data += [[_mark(key, row.get(key), cell) for key, _ in group["columns"]] for row in group["rows"]]
        table = Table(data, colWidths=widths, repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef0")), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                 ("VALIGN", (0, 0), (-1, -1), "TOP")]
        table.setStyle(TableStyle(style))
        story += [table, Spacer(1, 4)]
        if not group.get("summary_first"):
            story.append(Paragraph(" | ".join(f"<b>{_esc(label)}:</b> {_esc(value)}" for label, value in group["totals"]), styles["Normal"]))
        if group.get("layout") == "sheet":
            story += [Spacer(1, 16), KeepTogether(_signature(width))]
    if not (doc["groups"] and all(g.get("layout") == "sheet" for g in doc["groups"])):
        story += [Spacer(1, 22), KeepTogether(_signature(width))]
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
